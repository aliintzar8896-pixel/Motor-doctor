from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\HP\Desktop\Motor-Doctor")
APP = ROOT / "Motor-doctor-main"
OUT = ROOT / "Motor_Doctor_Project_Synopsis.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
styles['Normal'].font.size = Pt(12)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.15

for style_name, size in [('Title', 20), ('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)]:
    style = styles[style_name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(size)
    style.font.bold = True

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def para(text='', bold=False, align=None, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.first_line_indent = Inches(0.25) if text and align != WD_ALIGN_PARAGRAPH.CENTER else None
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def page_title(text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)

def bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def numbered(items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def diagram(filename, caption):
    path = ROOT / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.2))
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.runs[0].italic = True

# Cover
for _ in range(2):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('MOTOR DOCTOR'); r.bold = True; r.font.size = Pt(22); r.font.name = 'Times New Roman'
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Project Synopsis'); r.bold = True; r.font.size = Pt(18); r.font.name = 'Times New Roman'
for _ in range(2): doc.add_paragraph()
for text, size in [('BACHELOR OF COMPUTER APPLICATION', 14), ('5th SEMESTER', 14)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
doc.add_paragraph()
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(3.2); table.columns[1].width = Inches(3.2)
rows = [
    ('Project Guide:', 'Submitted By:'),
    ('Dr. Mithilesh Kumar Sharma\nMr. Yaduvansh Mani', 'Kashaf Shafeeq\nTCA2468249'),
    ('', 'SEC- E'),
    ('', '217360532258000'),
]
for i, row in enumerate(rows):
    for j, value in enumerate(row):
        set_cell_text(table.cell(i, j), value, bold=(i == 0 or i == 1))
doc.add_paragraph()
for text in ['FACULTY OF ENGINEERING & COMPUTING SCIENCES', 'TEERTHANKER MAHAVEER UNIVERSITY', 'MORADABAD-244001', 'SEPTEMBER, 2026']:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(14)

# Declaration
page_title('DECLARATION')
para('I, Kashaf Shafeeq (TCA2468249), student of BCA, 5th Semester, studying at Faculty of Engineering and Computer Sciences, Teerthanker Mahaveer University, Moradabad (U.P.), hereby declare that the Internship Report/Dissertation entitled “Motor Doctor” submitted in partial fulfilment for the degree of Bachelor of Computer Applications (BCA) is the original work carried out by me.')
para('The information and material presented in this synopsis are true to the best of my knowledge. This Internship Report/Dissertation has not been submitted to any other university or institution for the award of any degree, diploma or fellowship.')
para('Date: ____________')
para('Place: Moradabad')
para('Submitted By:\nKashaf Shafeeq (TCA2468249)', bold=True)

# Certificate
page_title('CERTIFICATE')
para('This is to certify that the Internship Report/Dissertation entitled “Motor Doctor” submitted in partial fulfilment of the requirements for the degree of Bachelor of Computer Applications (BCA) has been carried out by the student named below under my supervision and guidance.')
para('Submitted By:\nKashaf Shafeeq (TCA2468249)', bold=True)
para('Guide Name: Dr. Mithilesh Kumar Sharma')
para('Guide Signature: __________________________')
para('Date: ____________')

# Acknowledgement
page_title('ACKNOWLEDGEMENT')
para('I sincerely acknowledge the support and guidance of all the people who contributed to the successful completion of my internship project, “Motor Doctor”.')
para('First, I express my gratitude to my project guide, Dr. Mithilesh Kumar Sharma, and Mr. Yaduvansh Mani for their valuable guidance, suggestions and encouragement throughout the project.')
para('I also thank the Hon’ble Vice Chancellor, the Vice Principal, the Head of the Faculty of Engineering and Computer Sciences, and the Internship/Dissertation Coordinator of Teerthanker Mahaveer University for providing the academic environment and support required for this work.')
para('Finally, I thank my family and friends for their continuous encouragement and assistance during the development and documentation of this project.')
para('Date: ____________')
para('Kashaf Shafeeq (TCA2468249)', bold=True)

# Abstract
page_title('ABSTRACT')
para('Motor Doctor is a responsive web-based roadside assistance and vehicle service platform designed to help drivers handle vehicle breakdowns quickly and confidently. The application provides a single digital interface for emergency assistance, mechanic discovery, service request tracking, bill analysis and highway safety information.')
para('The frontend is developed using React with TypeScript and Vite. React Router is used for page navigation, Tailwind CSS is used for responsive styling, and Lucide React provides interface icons. The application includes role-oriented experiences for drivers, mechanics and administrators. Drivers can request assistance, select an issue and urgency level, view mechanic details, follow service status, choose a payment method and review safety guidance. Mechanics can view service requests through a dedicated portal, while administrators can access dashboard information.')
para('The Bill Doctor module helps users examine service quotations by comparing quoted prices with recommended fair prices and classifying items as essential, fair, overpriced or unnecessary. Emergency features include SOS access, contact information, service tracking and a map-based assistance view. The current implementation is a functional frontend prototype using structured mock data; live authentication, persistent database storage, GPS dispatch and payment gateway integration are proposed for future development.')
para('Keywords: Motor Doctor, Roadside Assistance, React, TypeScript, Emergency Vehicle Service, Mechanic Tracking, Bill Analysis, Web Application')

# Contents
page_title('TABLE OF CONTENTS')
contents = ['1. Project Title', '2. Domain', '3. Problem Statement', '4. Project Description', '4.1 Proposed System', '4.2 Need for the Project', '4.3 Objectives of the Project', '4.4 Aim of the Project', '4.5 Scope of the Project', '4.6 Background of Roadside Assistance', '4.7 Existing Vehicle Service Practices', '4.8 Digital Roadside Assistance Platforms', '5. Methodology', '5.1 Requirement Analysis', '5.2 System Design', '5.3 System Modules', '5.4 Development', '5.5 Testing', '5.6 Request-Response and Service Flow', '6. Results', '7. Future Scope', '8. Conclusion', '9. References']
for item in contents: para(item)

# Main content
page_title('1. PROJECT TITLE')
para('The title of the proposed project is “Motor Doctor: A Digital Roadside Assistance and Vehicle Service Platform”.')
heading('2. DOMAIN', 1)
para('Web Application Development using React, TypeScript and Vite, with a focus on roadside assistance, service request management, user experience and automotive support services.')
heading('3. PROBLEM STATEMENT', 1)
para('Vehicle breakdowns on roads and highways create safety, time and financial challenges for drivers. In many situations, a driver has difficulty finding a trusted mechanic nearby, estimating arrival time, communicating the vehicle problem, understanding repair charges and tracking the progress of a service request. Traditional assistance through phone calls and informal contacts is fragmented and provides limited visibility.')
para('The proposed Motor Doctor platform addresses these problems through a centralized interface that connects a driver with roadside assistance workflows, mechanic information, emergency contacts, service tracking and bill guidance. It aims to make the first response to a breakdown more organized, transparent and accessible.')

heading('4. PROJECT DESCRIPTION', 1)
heading('4.1 Proposed System', 2)
para('Motor Doctor is a responsive single-page web application that brings common roadside assistance tasks into one platform. The system is organized around driver, mechanic and administrator journeys. A driver can open the emergency assistance flow, enter vehicle and breakdown details, choose normal, urgent or SOS priority, view available mechanics and monitor the request status. The platform also includes a Bill Doctor service for reviewing repair quotations and a Highway Safety section for preventive guidance.')
para('The current prototype uses React components, typed domain models and structured mock data to demonstrate the complete user interface and service workflow. It is designed so that a production backend, authentication service, database, live location service and payment provider can be connected later without changing the overall user journey.')
heading('4.2 Need for the Project', 2)
bullets(['Provide a quick and organized way to request roadside assistance.', 'Help drivers identify suitable mechanics using availability, rating, distance and estimated arrival time.', 'Improve transparency in service quotations through bill analysis.', 'Support emergency communication with highway, police, ambulance and platform contacts.', 'Give mechanics a structured portal for viewing and managing incoming requests.', 'Demonstrate a practical, responsive and role-oriented automotive service application.'])
heading('4.3 Objectives of the Project', 2)
numbered(['To design and develop a responsive roadside assistance web application.', 'To create an emergency request workflow for different vehicle problems and urgency levels.', 'To present mechanic profiles with service specialties, rating, distance and estimated arrival time.', 'To provide a live-tracking style interface for the progress of a service request.', 'To develop a Bill Doctor module for comparing quoted and recommended repair prices.', 'To provide separate dashboard experiences for drivers, mechanics and administrators.', 'To improve awareness of highway safety and emergency support resources.', 'To gain practical experience in React, TypeScript, routing, component design and responsive UI development.'])
heading('4.4 Aim of the Project', 2)
para('The primary aim of Motor Doctor is to develop a user-friendly digital platform that helps motorists obtain roadside assistance, communicate vehicle problems, track service progress and make more informed decisions about repair quotations.')
heading('4.5 Scope of the Project', 2)
bullets(['Driver registration and login interface.', 'Emergency assistance and SOS request interface.', 'Vehicle, location, issue type, urgency and description capture.', 'Mechanic listing, filtering and mechanic profile presentation.', 'Request status tracking from pending to completed or cancelled.', 'Bill quotation analysis and expert recommendations.', 'Mechanic portal and administrator dashboard interfaces.', 'Payment method selection interface and service summary.', 'Highway safety guidance and emergency contact information.', 'Responsive browser-based experience for desktop and mobile screens.'])
para('The present scope is a frontend prototype. Persistent user accounts, live GPS tracking, real-time notifications, verified mechanic onboarding, cloud storage and actual payment processing are outside the current implementation and are included in the future scope.')
heading('4.6 Background of Roadside Assistance', 2)
para('Roadside assistance services support motorists when a vehicle becomes unusable because of a puncture, battery failure, engine problem, fuel issue, electrical fault, brake problem or accident. A digital platform can reduce response delays by collecting the relevant information at the first point of contact and presenting it to the service provider in a structured manner.')
heading('4.7 Existing Vehicle Service Practices', 2)
para('Existing roadside support commonly depends on phone calls, local references, workshop directories or separate service-provider applications. These methods may not provide a unified view of mechanic availability, estimated arrival time, request status and quotation fairness. Drivers may also have limited knowledge of which repair items are essential or optional.')
heading('4.8 Digital Roadside Assistance Platforms', 2)
para('A digital roadside assistance platform combines a user interface, service-request logic, mechanic information, location data and communication workflows. A typical flow is: Driver Interface -> Assistance Service -> Mechanic/Dispatch Interface -> Status Update -> Driver Notification. Motor Doctor models this flow through typed data structures and reusable React components.')

heading('5. METHODOLOGY', 1)
para('The project follows an incremental software development approach. Requirements were identified first, followed by interface design, component development, integration of page routes and validation of the major user workflows.')
heading('5.1 Requirement Analysis', 2)
heading('5.1.1 Functional Requirements', 3)
bullets(['The system shall provide navigation among home, emergency, bill analysis, safety, dashboard, portal, login and contact pages.', 'The system shall allow a user to create a roadside service request through a structured modal.', 'The system shall present available mechanics and their service details.', 'The system shall display request status and tracking information.', 'The system shall allow a user to analyse bill items and view recommendations.', 'The system shall provide role-specific dashboard and portal views.', 'The system shall provide payment method selection and service summary interfaces.', 'The system shall provide emergency contact and SOS access.'])
heading('5.1.2 Non-Functional Requirements', 3)
bullets(['Usability: interfaces should be understandable during stressful roadside situations.', 'Responsiveness: screens should adapt to desktop and mobile viewport sizes.', 'Maintainability: reusable components and TypeScript types should keep the code organized.', 'Performance: common navigation and interactions should feel immediate in the browser.', 'Accessibility: controls should have clear labels, visual hierarchy and keyboard-friendly structure.', 'Scalability: the frontend should support later connection to live APIs and a persistent backend.'])
heading('5.2 System Design', 2)
para('The application is designed as a component-based frontend with three logical layers: presentation, application state and service data. The presentation layer contains pages and reusable components. The application layer contains routing, context state and user interaction logic. The data layer currently contains typed mock data and is prepared for replacement by API calls and a database-backed service.')
para('Overall flow: User -> React Page/Component -> App Context and Domain Types -> Service/API Layer (future) -> Database and External Services (future) -> Updated Interface')
para('The main domain entities represented in the project are User, Mechanic, ServiceRequest, BillItem, BillAuditReport, Review, SOSContact and BookingNotification.')
diagram('arch_diagram_motor_doctor.png', 'Figure 1: Motor Doctor system architecture')
diagram('er_diagram_motor_doctor.png', 'Figure 2: Motor Doctor entity relationship design')
diagram('dfd_diagram_motor_doctor.png', 'Figure 3: Motor Doctor data flow diagram')
heading('5.3 System Modules', 2)
modules = [
('1. Home Page Module', 'Introduces the platform and directs users to roadside assistance, bill analysis, safety information and other services.'),
('2. Authentication Module', 'Provides login, signup and authentication-related screens for the driver and service-provider journeys.'),
('3. Emergency Assistance Module', 'Collects vehicle, issue, urgency and location details and starts an assistance request.'),
('4. SOS and Emergency Contact Module', 'Provides a prominent SOS flow and quick access to highway help, police, ambulance and Motor Doctor support.'),
('5. Mechanic Discovery Module', 'Displays mechanic cards with rating, review count, distance, ETA, service specialties, availability and towing support.'),
('6. Live Tracking Module', 'Presents service request progress, mechanic information, ETA, location and status updates in a tracking-oriented interface.'),
('7. Bill Doctor Module', 'Analyses quotation items, compares quoted and fair prices, identifies unnecessary or overpriced work and shows possible savings.'),
('8. Payment Module', 'Provides a payment method selector for cash, UPI options and payment-related service summary states.'),
('9. Mechanic Portal Module', 'Gives mechanics a dedicated view for service requests, customer details, vehicle problems and request actions.'),
('10. User Dashboard Module', 'Summarizes driver activity, active requests, service history and vehicle-related information.'),
('11. Admin Dashboard Module', 'Provides an administrative overview of users, mechanics, requests and platform activity.'),
('12. Highway Safety Module', 'Presents practical safety information and preventive guidance for highway travel and vehicle emergencies.'),
]
for title, text in modules:
    para(title, bold=True); para(text)
heading('5.4 Development', 2)
para('The project is developed in a local Windows environment using Visual Studio Code. The frontend uses React 18, TypeScript and Vite. React Router DOM manages routes, Tailwind CSS supports styling, Lucide React supplies icons, Recharts supports dashboard visualizations, and Sonner provides toast notifications. The source code is organized into pages, reusable components, context, library utilities and domain types.')
table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(['Software Requirement', 'Technology/Tool']): set_cell_text(table.cell(0, i), value, bold=True); shade(table.cell(0, i), 'D9EAF7')
for key, value in [('Operating System', 'Windows'), ('Frontend', 'React, TypeScript, Vite'), ('Styling', 'Tailwind CSS'), ('Routing', 'React Router DOM'), ('Icons and UI', 'Lucide React and Radix UI components'), ('Charts', 'Recharts'), ('Code Editor', 'Visual Studio Code'), ('Browser', 'Google Chrome / Microsoft Edge')]:
    cells = table.add_row().cells; set_cell_text(cells[0], key); set_cell_text(cells[1], value)
heading('5.5 Testing', 2)
para('Testing is performed through build validation, lint checks and manual interaction testing of the implemented frontend flows. The following areas are verified:')
bullets(['Route navigation and fallback routing.', 'Opening and closing emergency and consultation modals.', 'Form field entry and selection of issue and urgency values.', 'Mechanic cards, service details and request status presentation.', 'Bill analysis classifications and savings summary.', 'Payment method selection and toast notifications.', 'Responsive layout across desktop and mobile widths.', 'Dashboard and portal rendering using the available structured data.'])
heading('5.6 Request-Response and Service Flow', 2)
numbered(['The user opens Motor Doctor in a browser.', 'The user selects an assistance or information feature.', 'React Router renders the selected page.', 'The page reads the current application state and typed service data.', 'The user submits or changes a form, selection or action.', 'The application updates local context state and displays the resulting interface state.', 'In a production version, the same action will call an authenticated API and persist the result in a database.', 'The updated request, status or report is displayed to the user.'])

heading('6. RESULTS', 1)
para('The implemented prototype demonstrates the primary interface and navigation requirements of the Motor Doctor platform. The following results are available in the current build:')
bullets(['A responsive home page with clear entry points for emergency help, bill analysis and safety information.', 'A complete emergency request interface with vehicle, issue, urgency and location-oriented fields.', 'Mechanic cards showing service capabilities, ratings, distance, ETA and availability.', 'A tracking card that represents the progress of an active service request.', 'A Bill Doctor interface with item-level verdicts, reasons, fair-price comparison and potential savings.', 'Separate user, mechanic and administrator page experiences.', 'Authentication screens and protected-journey entry points.', 'Payment method selection, notifications and emergency contact presentation.', 'Architecture, entity relationship and data flow diagrams supporting the system design.'])
heading('7. FUTURE SCOPE', 1)
numbered(['Backend and Database Integration: Add a Node.js/Express or equivalent backend with PostgreSQL/Supabase for persistent users, mechanics, requests and bills.', 'Secure Authentication: Implement session or token-based authentication, password security and role-based access control.', 'Live GPS and Dispatch: Connect maps, geolocation, mechanic proximity search and real-time ETA updates.', 'Real-Time Notifications: Add SMS, email, push notifications and WebSocket-based status updates.', 'Payment Gateway: Integrate a secure production payment provider and transaction reconciliation.', 'Mechanic Verification: Add document verification, onboarding workflow, service zones and availability scheduling.', 'AI-Assisted Diagnosis: Use structured symptoms and optional image/audio input to suggest likely vehicle issues.', 'Digital Invoices and Reports: Generate downloadable invoices, service histories and bill audit reports.', 'Mobile Application: Extend the responsive web experience into Android and iOS applications.', 'Analytics and Monitoring: Add service response metrics, customer feedback analysis and operational dashboards.'])
heading('8. CONCLUSION', 1)
para('Motor Doctor presents a practical solution for organizing roadside assistance and vehicle service support through a modern web interface. The project combines emergency request handling, mechanic discovery, tracking, bill analysis, payment selection, safety guidance and role-oriented dashboards in one coherent application.')
para('The React and TypeScript implementation demonstrates component-based frontend development, client-side routing, structured domain modelling, responsive design and reusable interface patterns. Although the current version uses mock data and does not yet persist transactions in a backend, its architecture establishes a clear foundation for live dispatch, authentication, database storage, payments and notifications. The project therefore fulfils its purpose as an internship-level prototype while offering a realistic path toward a production roadside assistance platform.')
heading('9. REFERENCES', 1)
refs = ['React Documentation, Meta Open Source.', 'TypeScript Documentation, Microsoft.', 'Vite Documentation, Vite Team.', 'Tailwind CSS Documentation.', 'React Router Documentation.', 'Lucide React Documentation.', 'Recharts Documentation.', 'MDN Web Docs, HTML, CSS and JavaScript Web Standards.']
for ref in refs: para(ref)

doc.save(OUT)
print(f'Created: {OUT}')